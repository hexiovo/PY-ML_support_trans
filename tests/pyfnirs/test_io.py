"""Output fidelity, atomic publication, and no-overwrite guarantees."""

from __future__ import annotations

import csv
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from pyfnirs_converter.writers import write_csv_bundle, write_xlsx


def _typed(text: str) -> dict[str, object]:
    return {"matlab_class": "string", "shape": [1, 1], "value": text, "is_missing": False}


def _payload() -> dict[str, object]:
    return {
        "schema_version": 1,
        "X": [[1.25, 3.75], [2.5, float("nan")]],
        "FeatureIDs": ["F-O2", "F-HbR"],
        "Samples": [{"ObservationID": _typed("OBS-42")}, {"ObservationID": _typed("OBS-8")}],
        "Targets": {"Group": ["control", "case"]},
        "Covariates": {"AgeYears": [31.0, 45.0]},
        "FeatureDefinitions": [
            {"FeatureID": _typed("F-O2"), "Metric": "synthetic_mean"},
            {"FeatureID": _typed("F-HbR"), "Metric": "synthetic_mean"},
        ],
        "ValidMask": [[True, True], [True, False]],
        "MissingReasons": [["", ""], ["", "qc_rejected"]],
        "Provenance": {
            "valid_mask": [[True, True], [True, False]],
            "missing_reasons": [["", ""], ["", "qc_rejected"]],
            "synthetic": True,
        },
    }


class ConverterWriterTest(unittest.TestCase):
    def test_csv_and_xlsx_preserve_ids_labels_values_and_missing_reason(self) -> None:
        with tempfile.TemporaryDirectory(prefix="pyfnirs-writer-中文 ") as temporary:
            root = Path(temporary)
            csv_dir = root / "CSV 输出"
            paths = write_csv_bundle(_payload(), csv_dir)
            self.assertEqual(
                [path.name for path in paths],
                ["FeatureMatrix.csv", "Samples.csv", "FeatureDictionary.csv", "Provenance.json"],
            )
            with (csv_dir / "FeatureMatrix.csv").open(encoding="utf-8-sig", newline="") as stream:
                matrix = list(csv.reader(stream))
            self.assertEqual(matrix, [
                ["ObservationID", "Group", "AgeYears", "F-O2", "F-HbR"],
                ["OBS-42", "control", "31", "1.25", "3.75"],
                ["OBS-8", "case", "45", "2.5", ""],
            ])
            with (csv_dir / "FeatureDictionary.csv").open(encoding="utf-8-sig", newline="") as stream:
                dictionary = list(csv.reader(stream))
            self.assertEqual([row[0] for row in dictionary[1:]], ["F-O2", "F-HbR"])
            provenance = json.loads((csv_dir / "Provenance.json").read_text(encoding="utf-8"))
            self.assertEqual(provenance["missing_reasons"], [["", ""], ["", "qc_rejected"]])

            workbook_path = root / "review.xlsx"
            write_xlsx(_payload(), workbook_path)
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    ["FeatureMatrix", "Samples", "FeatureDictionary", "Provenance"],
                )
                rows = list(workbook["FeatureMatrix"].values)
                self.assertEqual(rows[0], ("ObservationID", "Group", "AgeYears", "F-O2", "F-HbR"))
                self.assertEqual(rows[1], ("OBS-42", "control", 31, 1.25, 3.75))
                self.assertEqual(rows[2], ("OBS-8", "case", 45, 2.5, None))
                provenance_rows = list(workbook["Provenance"].values)
                provenance_values = {row[0]: row[1] for row in provenance_rows[1:]}
                self.assertEqual(json.loads(provenance_values["missing_reasons"]), [["", ""], ["", "qc_rejected"]])
            finally:
                workbook.close()

    def test_existing_outputs_are_never_overwritten(self) -> None:
        with tempfile.TemporaryDirectory(prefix="pyfnirs-collision-") as temporary:
            root = Path(temporary)
            csv_dir = root / "existing-bundle"
            csv_dir.mkdir()
            marker = csv_dir / "user.txt"
            marker.write_text("keep", encoding="utf-8")
            with self.assertRaises(FileExistsError):
                write_csv_bundle(_payload(), csv_dir)
            self.assertEqual(marker.read_text(encoding="utf-8"), "keep")

            xlsx = root / "existing.xlsx"
            xlsx.write_bytes(b"user-owned file")
            before = xlsx.read_bytes()
            with self.assertRaises(FileExistsError):
                write_xlsx(_payload(), xlsx)
            self.assertEqual(xlsx.read_bytes(), before)

    def test_equals_prefixed_strings_remain_literal_text_in_csv_and_xlsx(self) -> None:
        payload = _payload()
        payload["FeatureIDs"] = ["=F-O2", "F-HbR"]
        payload["Samples"] = [
            {"ObservationID": _typed("=OBS-42")},
            {"ObservationID": _typed("OBS-8")},
        ]
        payload["Targets"] = {"Group": ["=control", "case"]}
        payload["FeatureDefinitions"] = [
            {"FeatureID": _typed("=F-O2"), "Metric": "synthetic_mean"},
            {"FeatureID": _typed("F-HbR"), "Metric": "synthetic_mean"},
        ]
        provenance = dict(payload["Provenance"])
        provenance["source_record_id"] = "=source-42"
        payload["Provenance"] = provenance

        with tempfile.TemporaryDirectory(prefix="pyfnirs-literal-text-") as temporary:
            root = Path(temporary)
            csv_dir = root / "literal-csv"
            write_csv_bundle(payload, csv_dir)
            with (csv_dir / "FeatureMatrix.csv").open(encoding="utf-8-sig", newline="") as stream:
                matrix = list(csv.reader(stream))
            self.assertEqual(matrix[0][3], "=F-O2")
            self.assertEqual(matrix[1][:2], ["=OBS-42", "=control"])

            workbook_path = root / "literal.xlsx"
            write_xlsx(payload, workbook_path)
            formula_view = load_workbook(workbook_path, read_only=True, data_only=False)
            value_view = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                matrix_sheet = formula_view["FeatureMatrix"]
                self.assertEqual(matrix_sheet["D1"].value, "=F-O2")
                self.assertEqual(matrix_sheet["D1"].data_type, "s")
                self.assertEqual(matrix_sheet["A2"].value, "=OBS-42")
                self.assertEqual(matrix_sheet["A2"].data_type, "s")
                self.assertEqual(matrix_sheet["B2"].value, "=control")
                self.assertEqual(matrix_sheet["B2"].data_type, "s")
                self.assertEqual(value_view["FeatureMatrix"]["A2"].value, "=OBS-42")
                self.assertEqual(value_view["FeatureMatrix"]["B2"].value, "=control")
                dictionary_cell = formula_view["FeatureDictionary"]["A2"]
                self.assertEqual(dictionary_cell.value, "=F-O2")
                self.assertEqual(dictionary_cell.data_type, "s")
                provenance_rows = list(value_view["Provenance"].values)
                provenance_values = {row[0]: row[1] for row in provenance_rows[1:]}
                self.assertEqual(json.loads(provenance_values["source_record_id"]), "=source-42")
                source_value_cell = next(
                    row[1]
                    for row in formula_view["Provenance"].iter_rows()
                    if row[0].value == "source_record_id"
                )
                self.assertEqual(source_value_cell.data_type, "s")
            finally:
                formula_view.close()
                value_view.close()

    def test_xlsx_preserves_finite_double_binary64_values_as_numeric(self) -> None:
        payload = _payload()
        source_values = [
            1.2345678901234567,
            -0.0,
            float.fromhex("0x1.0000000000001p+0"),
            float.fromhex("0x0.0000000000001p-1022"),
            float.fromhex("0x1.fffffffffffffp+1023"),
        ]
        feature_ids = [f"F-{index}" for index in range(len(source_values))]
        payload["FeatureIDs"] = feature_ids
        payload["X"] = [source_values.copy(), source_values.copy()]
        payload["ValidMask"] = [[True] * len(source_values) for _ in range(2)]
        payload["MissingReasons"] = [[""] * len(source_values) for _ in range(2)]
        payload["FeatureDefinitions"] = [
            {"FeatureID": _typed(feature_id), "Metric": "synthetic_mean"}
            for feature_id in feature_ids
        ]
        provenance = dict(payload["Provenance"])
        provenance["valid_mask"] = payload["ValidMask"]
        provenance["missing_reasons"] = payload["MissingReasons"]
        payload["Provenance"] = provenance

        with tempfile.TemporaryDirectory(prefix="pyfnirs-double-precision-") as temporary:
            root = Path(temporary)
            csv_dir = root / "precision-csv"
            write_csv_bundle(payload, csv_dir)
            with (csv_dir / "FeatureMatrix.csv").open(encoding="utf-8-sig", newline="") as stream:
                csv_rows = list(csv.DictReader(stream))
            for feature_id, source_value in zip(feature_ids, source_values):
                literal = csv_rows[0][feature_id]
                self.assertEqual(float(literal).hex(), source_value.hex())

            workbook_path = root / "precision.xlsx"
            write_xlsx(payload, workbook_path)
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                for column, source_value in enumerate(source_values, start=4):
                    cell = workbook["FeatureMatrix"].cell(row=2, column=column)
                    self.assertEqual(cell.data_type, "n")
                    self.assertEqual(float(cell.value).hex(), source_value.hex())
            finally:
                workbook.close()

    def test_failed_csv_staging_does_not_publish_a_partial_bundle(self) -> None:
        with tempfile.TemporaryDirectory(prefix="pyfnirs-clean-stage-") as temporary:
            root = Path(temporary)
            payload = _payload()
            payload["Provenance"] = {"unsupported": object()}
            destination = root / "failed-bundle"
            with self.assertRaises(TypeError):
                write_csv_bundle(payload, destination)
            self.assertFalse(destination.exists())
            self.assertEqual(list(root.glob(".failed-bundle-*")), [])


if __name__ == "__main__":
    unittest.main()
