"""Lossless table reshaping, explicit identities and isolated batch publication."""
from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import shutil
import subprocess
import tempfile
import threading
import uuid
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import __version__


class ConversionError(ValueError):
    pass


class Cancelled(ConversionError):
    pass


class Skipped(ConversionError):
    pass


@dataclass(frozen=True)
class Options:
    sample_columns: tuple[str, ...] = ()
    coordinate_columns: tuple[str, ...] = ()
    value_columns: tuple[str, ...] = ()
    table_path: str = ""
    sheet: str = ""
    matlab: str = ""
    recursive: bool = True
    merge: bool = True
    metadata: str = ""
    metadata_sheet: str = ""
    join_columns: tuple[str, ...] = ()
    target: str = ""
    metadata_only: bool = False


@dataclass
class Table:
    frame: pd.DataFrame
    origin: str
    role: str = "feature"
    native: dict | None = None


@dataclass
class Converted:
    samples: list[dict]
    features: dict[str, dict]
    cells: dict[tuple[str, str], str]
    tables: list[Table]
    sample_columns: tuple[str, ...]


def check_cancel(cancel: threading.Event | None):
    if cancel is not None and cancel.is_set():
        raise Cancelled("已取消；已完成的文件保留，未完成文件不发布。")


def digest_file(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def token(value) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def read_csv(path: Path) -> pd.DataFrame:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    if not rows or not rows[0] or any(not name for name in rows[0]):
        raise ConversionError("CSV 表头为空。")
    if len(set(rows[0])) != len(rows[0]):
        raise ConversionError("CSV 列名重复。")
    if any(len(row) != len(rows[0]) for row in rows[1:]):
        raise ConversionError("CSV 存在列数不一致的行。")
    return pd.DataFrame(rows[1:], columns=rows[0], dtype=object)


def read_xlsx(path: Path, selected: str = "") -> list[Table]:
    result = []
    # Cached formula values must exist; formulas are never silently turned into blanks.
    with path.open("rb") as source:
        values = load_workbook(source, read_only=True, data_only=True)
        formulas = load_workbook(path, read_only=True, data_only=False)
        try:
            if selected and selected not in values.sheetnames:
                raise ConversionError(f"没有工作表 {selected!r}；可选：{values.sheetnames}")
            for name in ([selected] if selected else values.sheetnames):
                rows = list(values[name].iter_rows(values_only=True))
                raw = list(formulas[name].iter_rows())
                if not rows:
                    continue
                for i, row in enumerate(raw):
                    for j, cell in enumerate(row):
                        if cell.data_type == "f" and rows[i][j] is None:
                            raise ConversionError(f"{name}!{cell.coordinate} 公式没有缓存值；请用 Excel 计算并保存。")
                header = [text(item) for item in rows[0]]
                invalid_header = not header or any(not item for item in header) or len(set(header)) != len(header)
                grid = None
                if invalid_header:
                    if name == "分析长表":
                        raise ConversionError(f"{name} 的表头为空或重复。")
                    # Reference README pages are cell grids, not training tables.
                    # Preserve even the first row under explicit Excel coordinates.
                    header = ["__excel_column_" + get_column_letter(i + 1) for i in range(max(len(row) for row in rows))]
                    body = rows
                    grid = {"excel_grid": True, "header_policy": "保留所有单元格；第一行仍为数据，列名是原 Excel 列坐标。"}
                else:
                    body = rows[1:]
                if any(len(row) > len(header) for row in body):
                    raise ConversionError(f"{name} 有单元格超出表头范围。")
                frame = pd.DataFrame([[text(row[i]) if i < len(row) else "" for i in range(len(header))] for row in body], columns=header, dtype=object)
                result.append(Table(frame, f"{path}#{name}", "feature" if name == "分析长表" else "metadata", grid))
        finally:
            values.close()
            formulas.close()
    return result


def find_matlab(selected: str) -> str:
    if selected:
        path = Path(selected).expanduser().resolve()
        if not path.is_file() or path.name.casefold() not in {"matlab.exe", "matlab"}:
            raise ConversionError("请选择实际的 matlab.exe。")
        return str(path)
    found = shutil.which("matlab")
    if not found:
        raise ConversionError("MAT 需要已安装并有许可的 MATLAB；请指定 matlab.exe，或先从 MATLAB 导出特征长表 CSV。")
    return found


def native_mat(path: Path, options: Options, cancel=None) -> list[Table]:
    bridge = Path(__file__).resolve().parent / "matlab"
    with tempfile.TemporaryDirectory(prefix="pyml-hypereeg-") as temporary:
        root = Path(temporary)
        config = root / "request.json"
        config.write_text(token({"input": str(path.resolve()), "output": str(root), "table_path": options.table_path}), encoding="utf-8")
        quoted = lambda item: str(item).replace("'", "''")
        expression = f"addpath('{quoted(bridge)}'); pyml_export_table('{quoted(config)}')"
        log_path = root / "matlab.log"
        with log_path.open("wb") as log:
            process = subprocess.Popen(
                [find_matlab(options.matlab), "-wait", "-batch", expression], stdout=log, stderr=log,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
            )
            try:
                while process.poll() is None:
                    if cancel is not None and cancel.wait(0.15):
                        if os.name == "nt":
                            subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True)
                        else:
                            process.terminate()
                        process.wait()
                        check_cancel(cancel)
                    if cancel is None:
                        threading.Event().wait(0.15)
            finally:
                if process.poll() is None:
                    process.terminate()
                    process.wait()
        if process.returncode:
            message = log_path.read_text(encoding="utf-8", errors="replace")[-4000:]
            raise ConversionError(f"MATLAB 读取失败：{message}")
        native = json.loads((root / "table.json").read_text(encoding="utf-8-sig"))
        frame = read_csv(root / "table.csv")
        quality = native.get("quality", {})
        if quality.get("status") in {"failed", "skipped", "unavailable", "dependency_blocked"}:
            raise Skipped("产物自身质量状态不合格：" + str(quality))
        if native.get("scopeKey"):
            if "__scope_key" in frame:
                raise ConversionError("载荷使用了保留列 __scope_key。")
            frame["__scope_key"] = text(native["scopeKey"])
        source_ids = native.get("sourceIds", [])
        if isinstance(source_ids, str):
            source_ids = [source_ids]
        if native.get("scope") == "record" and len(source_ids) == 1:
            if "__source_id" in frame:
                raise ConversionError("载荷使用了保留列 __source_id。")
            frame["__source_id"] = text(source_ids[0])
        return [Table(frame, f"{path}#{native['path']}", "feature", native)]


def load_tables(path: Path, options: Options, cancel=None) -> list[Table]:
    check_cancel(cancel)
    if path.suffix.casefold() == ".csv":
        return [Table(read_csv(path), str(path))]
    if path.suffix.casefold() == ".xlsx":
        tables = read_xlsx(path, options.sheet)
        # Current formal workbooks contain paths, not numerical observations.
        indexes = [item for item in tables if item.origin.endswith("#特征文件清单")]
        if indexes and not options.metadata_only:
            output = list(tables)
            seen = set()
            for index in indexes:
                if "path" not in index.frame:
                    raise ConversionError("特征文件清单缺少当前源码规定的 path 列。")
                for row in index.frame.to_dict("records"):
                    check_cancel(cancel)
                    if row.get("status") not in {"completed", "complete", "reused"}:
                        continue
                    target = Path(row["path"])
                    if not target.is_absolute():
                        target = path.parent / target
                    target = target.resolve()
                    if target in seen:
                        raise ConversionError("特征文件清单重复引用同一个 MAT。")
                    if not target.is_file() or target.suffix.casefold() != ".mat":
                        raise ConversionError(f"清单引用的正式 MAT 不存在：{target}")
                    seen.add(target)
                    output.extend(native_mat(target, options, cancel))
            return output
        return tables
    if path.suffix.casefold() == ".mat":
        if path.name.endswith(".invalid.mat") or path.name.endswith(".status.mat"):
            raise Skipped("状态/失效标记不是特征载荷。")
        return native_mat(path, options, cancel)
    raise Skipped("只支持正式 MAT、CSV、XLSX 特征/参考表。")


LONG_SAMPLES = ("file_name", "subject_id", "session", "condition", "timepoint", "segment_name", "window_index")
SINGLE_SAMPLES = ("StreamID", "SourcePath", "SubjectId", "IndependentUnitID", "Session", "Condition", "SourceSegmentId", "TimeSegmentIndex", "WindowIndex")
PAIR_SAMPLES = ("Session", "GroupId", "ConditionA", "ConditionB", "SubjectA", "SubjectB", "IndependentUnitIDA", "IndependentUnitIDB", "StreamIDA", "StreamIDB", "SourceSegmentIDA", "SourceSegmentIDB", "WindowIndex")
SINGLE_METRICS = ("Mean", "Std", "RMS", "PeakToPeak", "HjorthActivity", "HjorthMobility", "HjorthComplexity", "SampleEntropy", "ApproximateEntropy", "PermutationEntropy", "SpectralEntropy")
RECORD_METRICS = (
    "std", "variance", "rms", "mad", "iqr", "skewness", "kurtosis", "mean_abs_diff",
    "line_length", "zero_cross_rate", "hjorth_activity", "hjorth_mobility", "hjorth_complexity",
    "mean", "median", "trimmed_mean", "min", "max", "range", "peak_to_peak", "energy", "sum_abs",
    "total_power", "theta_beta_ratio", "alpha_theta_ratio", "alpha_beta_over_delta_theta",
    "peak_frequency_hz", "peak_power", "spectral_centroid_hz", "median_frequency_hz", "bandwidth_hz",
    "spectral_flatness", "individual_alpha_peak_hz", "individual_alpha_peak_power", "aperiodic_exponent",
    "aperiodic_offset", "faa", "normalized_spectral_entropy", "differential_entropy", "sample_entropy",
    "approximate_entropy", "fuzzy_entropy", "permutation_entropy", "svd_entropy", "lz_complexity",
    "lz_complexity_raw", "hurst_exponent", "dfa_alpha", "higuchi_fd", "petrosian_fd", "katz_fd",
    "mse_mean", "mse_complexity_index", "correlation_dimension", "largest_lyapunov_exponent",
)


def present(frame, candidates):
    return tuple(name for name in candidates if name in frame.columns)


def schema(table: Table, options: Options):
    frame = table.frame
    columns = set(frame.columns)
    if {"file_name", "feature", "value"} <= columns:
        samples, coordinates, values = present(frame, LONG_SAMPLES), ("feature",), ("value",)
    elif {"result_key", "measure", "cell_key", "value"} <= columns:
        samples = present(frame, LONG_SAMPLES)
        if "file_name" not in samples:
            samples = present(frame, ("source_mat", "subject_id", "session", "condition", "timepoint", "segment_name", "window_index"))
        coordinates, values = ("result_key", "measure", "cell_key"), ("value",)
    elif {"SubjectA", "SubjectB", "Metric", "Value"} <= columns:
        samples, coordinates, values = present(frame, PAIR_SAMPLES), present(frame, ("Method", "Metric", "Band")), present(frame, ("Value", "LagSeconds"))
    elif {"StreamID", "Band", "Power", "RelativePower"} <= columns:
        samples, coordinates, values = present(frame, SINGLE_SAMPLES), ("Band", "LowHz", "HighHz"), ("Power", "RelativePower")
    elif {"StreamID", "SubjectId", "WindowIndex"} <= columns:
        samples, coordinates, values = present(frame, SINGLE_SAMPLES), (), present(frame, SINGLE_METRICS)
    elif {"IndependentUnitID", "GlobalEfficiency", "Method", "Band"} <= columns:
        samples = present(frame, ("Session", "GroupId", "IndependentUnitID", "Condition", "WindowIndex"))
        coordinates, values = ("Method", "Band"), present(frame, ("Density", "MeanStrength", "MeanClustering", "GlobalEfficiency"))
    elif "__source_id" in columns and table.native and "channel_index" in columns:
        samples = ("__source_id",) + present(frame, ("window_index",))
        coordinates = present(frame, ("channel_index", "channel_name", "band", "band_low_hz", "band_high_hz"))
        values = tuple(name for name in frame.columns if name in RECORD_METRICS or name.startswith(("quantile_", "absolute_power_", "relative_power_", "log_power_db_", "log_power_log10_", "spectral_edge_", "rqa_")))
    elif table.role == "metadata" and not options.value_columns:
        raise Skipped("参考/分组/审计表保留为附表，不作为训练特征。")
    else:
        samples, coordinates, values = (), (), ()
    if options.sample_columns:
        samples = options.sample_columns
    if options.coordinate_columns:
        coordinates = options.coordinate_columns
    if options.value_columns:
        values = options.value_columns
    if not samples or not values:
        raise ConversionError(f"无法安全确定样本或数值列，请填写映射。实际列：{', '.join(frame.columns)}")
    all_names = samples + coordinates + values
    if len(set(all_names)) != len(all_names) or any(name not in columns for name in all_names):
        raise ConversionError("映射列缺失、重复或角色重叠。")
    return samples, coordinates, values


def numeric(value: str):
    if value.casefold() in {"true", "false"}:
        return True
    if value == "" or value.casefold() in {"nan", "<missing>"}:
        return False
    try:
        return math.isfinite(float(value))
    except (ValueError, OverflowError) as exc:
        raise ConversionError(f"数值列包含非数值 {value!r}，请修正映射；不会自动填补。") from exc


def convert_tables(tables: list[Table], options: Options, cancel=None) -> Converted:
    cells, features, samples = {}, {}, {}
    expected = None
    for table in tables:
        check_cancel(cancel)
        if options.metadata_only:
            continue
        try:
            sample_columns, coordinates, values = schema(table, options)
        except Skipped:
            continue
        if expected is not None and expected != sample_columns:
            raise ConversionError("不同表使用不同样本键；请分别转换或显式指定相同样本列。")
        expected = sample_columns
        if table.frame.empty:
            continue
        for row_number, row in enumerate(table.frame.to_dict("records"), start=2):
            if row_number % 256 == 0:
                check_cancel(cancel)
            # Retain QA/status as metadata. No unexplained row deletion.
            identity = [[name, row[name]] for name in sample_columns]
            if any(not value for _, value in identity):
                raise ConversionError(f"第 {row_number} 行样本键为空：{sample_columns}")
            key = token(identity)
            keep = set(sample_columns) | {"group", "dyad_id", "pair_id", "cluster_id", "age", "sex", "covariate_score", "GroupId", "RawGroup", "RawUser", "RawDevice", "Key", "WindowStart_s", "WindowEnd_s", "enabled", "Status", "ReviewStatus", "MappingStatus", "TimebaseStatus"}
            if options.target:
                keep.add(options.target)
            metadata = {f"meta__{name}": value for name, value in row.items() if name in keep and name not in values + coordinates}
            # All other columns, including per-feature QA, stay in source_tables.
            status = row.get("Status", row.get("status", "")).casefold()
            if status in {"failed", "skipped", "unavailable", "invalid", "rejected", "dependency_blocked"}:
                raise ConversionError(f"第 {row_number} 行状态为 {status}，没有进入训练资格；请先核查源项目质量结果。")
            if key in samples:
                for name, value in metadata.items():
                    if name in samples[key] and samples[key][name] != value:
                        raise ConversionError(f"同一样本的 {name} 不一致；不能自动平均/删行。")
                samples[key].update(metadata)
            else:
                samples[key] = {"sample_id": hashlib.sha256(key.encode("utf-8")).hexdigest(), **metadata, "_key": key}
            for value_column in values:
                value = row[value_column]
                numeric(value)
                coordinate = [[name, row[name]] for name in coordinates]
                # Stable, collision-free column names independent of source file order.
                feature_identity = {"values": value_column, "coordinates": coordinate}
                if table.native and "path" in table.native and not {"result_key", "measure", "cell_key"} <= set(row):
                    feature_identity["native_path"] = table.native["path"]
                    if table.native.get("producerId"):
                        feature_identity["producer_id"] = table.native["producerId"]
                column = "feature__" + token(feature_identity)
                if (key, column) in cells:
                    raise ConversionError(f"重复样本×特征（第 {row_number} 行）：{column}；没有自动聚合。")
                cells[key, column] = value
                record = features.setdefault(column, {"column": column, **feature_identity, "origins": []})
                if table.origin not in record["origins"]:
                    record["origins"].append(table.origin)
    if not cells and not options.metadata_only:
        raise Skipped("没有训练特征；参考表可用“只导出参考工作表”保存。")
    return Converted(list(samples.values()), features, cells, tables, expected or ())


def merge_converted(items: list[Converted]) -> Converted:
    merged = Converted([], {}, {}, [], ())
    samples = {}
    for item in items:
        if not item.cells:
            continue
        if merged.sample_columns and merged.sample_columns != item.sample_columns:
            raise ConversionError("合并失败：文件样本键不同，请分批转换或显式指定同一套映射。")
        merged.sample_columns = item.sample_columns
        for sample in item.samples:
            key = sample["_key"]
            if key in samples:
                for name, value in sample.items():
                    if name in samples[key] and samples[key][name] != value:
                        raise ConversionError(f"合并失败：同一样本的 {name} 冲突。")
                samples[key].update(sample)
            else:
                samples[key] = sample.copy()
        for key, value in item.cells.items():
            if key in merged.cells:
                raise ConversionError("合并失败：相同样本×特征重复；请排除重复文件/清单或分别输出。")
            merged.cells[key] = value
        for column, dictionary in item.features.items():
            if column in merged.features:
                merged.features[column]["origins"] = list(dict.fromkeys(merged.features[column]["origins"] + dictionary["origins"]))
            else:
                merged.features[column] = {**dictionary, "origins": list(dictionary["origins"])}
        merged.tables.extend(item.tables)
    merged.samples = list(samples.values())
    return merged


def attach_metadata(item: Converted, options: Options):
    if not options.metadata:
        if options.target:
            for sample in item.samples:
                source = "meta__" + options.target
                if "target__" + options.target in sample:
                    continue
                if source not in sample or not sample[source]:
                    raise ConversionError(f"指定标签 {options.target!r} 缺失；不会截断或伪造。")
                sample["target__" + options.target] = sample.pop(source)
        return
    if not options.join_columns:
        raise ConversionError("外部标签/分组表必须明确填写连接列，按完整原值连接，不猜文件名。")
    path = Path(options.metadata).expanduser().resolve()
    tables = load_tables(path, Options(sheet=options.metadata_sheet, metadata_only=True))
    matches = [table for table in tables if set(options.join_columns) <= set(table.frame.columns)]
    if len(matches) != 1:
        raise ConversionError("元数据连接表不存在或多个工作表都匹配；请通过工作表名明确选择。")
    rows = {}
    for row in matches[0].frame.to_dict("records"):
        key = tuple(row[name] for name in options.join_columns)
        if any(not value for value in key) or key in rows:
            raise ConversionError("元数据连接键为空或重复，无法进行多对一连接。")
        rows[key] = row
    for sample in item.samples:
        try:
            key = tuple(sample["meta__" + name] for name in options.join_columns)
            row = rows[key]
        except KeyError as exc:
            raise ConversionError("元数据存在未匹配样本；不会丢弃样本或截断标签。") from exc
        for name, value in row.items():
            column = "target__" + name if name == options.target else "meta__" + name
            if column in sample and sample[column] != value:
                raise ConversionError(f"元数据冲突：{name}")
            sample[column] = value
        if options.target and (options.target not in row or not row[options.target]):
            raise ConversionError("指定标签列不存在或标签缺失。")
        if options.target:
            sample.pop("meta__" + options.target, None)


def dataset_frame(item: Converted) -> pd.DataFrame:
    columns = list(dict.fromkeys(name for sample in item.samples for name in sample if name != "_key"))
    columns += list(item.features)
    rows = []
    for sample in item.samples:
        key = sample["_key"]
        rows.append({**{name: value for name, value in sample.items() if name != "_key"}, **{name: item.cells.get((key, name), "") for name in item.features}})
    return pd.DataFrame(rows, columns=columns)


def write_csv(frame: pd.DataFrame, path: Path):
    frame.to_csv(path, index=False, encoding="utf-8-sig", lineterminator="\n")


def write_package(item: Converted, destination: Path, options: Options, source: Path | None, cancel=None):
    check_cancel(cancel)
    destination.mkdir()
    raw = destination / "source_tables"
    raw.mkdir()
    raw_files = []
    for number, table in enumerate(item.tables, start=1):
        check_cancel(cancel)
        filename = f"table_{number:04d}.csv"
        write_csv(table.frame, raw / filename)
        origin_file = Path(table.origin.split("#", 1)[0])
        raw_files.append({"file": "source_tables/" + filename, "origin": table.origin, "role": table.role, "shape": list(table.frame.shape), "native": table.native, "source_sha256": digest_file(origin_file) if origin_file.is_file() else None})
    if item.cells:
        attach_metadata(item, options)
        frame = dataset_frame(item)
        write_csv(frame, destination / "dataset.csv")
        mask = pd.DataFrame([{ "sample_id": sample["sample_id"], **{column: int(numeric(item.cells.get((sample["_key"], column), ""))) for column in item.features}} for sample in item.samples])
        write_csv(mask, destination / "valid_mask.csv")
        (destination / "feature_dictionary.json").write_text(json.dumps(list(item.features.values()), ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {"schema_version": 1, "converter": "HyperEEG", "version": __version__, "created_at": datetime.now().astimezone().isoformat(), "options": asdict(options), "source": str(source) if source else "merged", "source_sha256": digest_file(source) if source else None, "rows": len(item.samples), "features": len(item.features), "sample_columns": item.sample_columns, "tables": raw_files, "feature_columns": list(item.features), "target_column": "target__" + options.target if options.target else None, "metadata_source_sha256": digest_file(Path(options.metadata)) if options.metadata else None, "qa_policy": "原始状态/质量列保留为元数据；valid_mask=1仅表示数值有限。训练前须按真实资格筛选。", "missing_policy": "保留空值/NaN/Inf及掩码；未插补、归一化、重算或拟合。"}
    (destination / "provenance.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    check_cancel(cancel)


def scan(source: Path, output: Path, options: Options) -> list[Path]:
    source, output = source.resolve(), output.resolve()
    if not source.is_dir():
        raise ConversionError("输入目录不存在。")
    if source == output or source.is_relative_to(output):
        raise ConversionError("输出目录不能与输入目录相同或包含输入目录。")
    candidates = source.rglob("*") if options.recursive else source.glob("*")
    found = []
    for path in sorted(candidates):
        if not path.is_file() or path.suffix.casefold() not in {".csv", ".xlsx", ".mat"}:
            continue
        resolved = path.resolve()
        if resolved.is_relative_to(output):
            continue
        # A published package remains an output even after it has been moved.
        if any((parent / "provenance.json").exists() or (parent / "batch_summary.json").exists() for parent in path.parents if parent.is_relative_to(source)):
            continue
        if options.metadata and resolved == Path(options.metadata).resolve():
            continue
        found.append(resolved)
    if not options.metadata_only:
        indexed = set()
        for path in found:
            if path.suffix.casefold() != ".xlsx":
                continue
            try:
                tables = read_xlsx(path, "特征文件清单")
            except Exception:
                # Invalid workbooks are diagnosed per file during conversion.
                continue
            for table in tables:
                if "path" not in table.frame:
                    continue
                for row in table.frame.to_dict("records"):
                    if row.get("status") in {"completed", "complete", "reused"} and row.get("path"):
                        referenced = Path(row["path"])
                        if not referenced.is_absolute():
                            referenced = path.parent / referenced
                        indexed.add(referenced.resolve())
        # A formal workbook consumes its referenced MATs once. Scanning those
        # same files again would duplicate every sample × feature during merge.
        found = [path for path in found if path not in indexed]
    return found


def preview(path: Path, options: Options, cancel=None):
    item = convert_tables(load_tables(path, options, cancel), options, cancel)
    attach_metadata(item, options)
    return item, (item.tables[0].frame if options.metadata_only and item.tables else dataset_frame(item)).head(30)


def batch(source: Path, output: Path, options: Options, cancel=None, progress: Callable | None = None) -> dict:
    files = scan(source, output, options)
    if not files:
        raise ConversionError("没有可转换输入（MAT/CSV/XLSX），请检查目录/递归选项。")
    check_cancel(cancel)
    output.mkdir(parents=True, exist_ok=True)
    run = output / ("HyperEEG_" + datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:8])
    run.mkdir()
    records, items = [], []
    summary = {"run_directory": str(run), "total": len(files), "success": 0, "failed": 0, "skipped": 0, "cancelled": False, "merge": None, "files": records}
    for number, path in enumerate(files, start=1):
        if cancel is not None and cancel.is_set():
            summary["cancelled"] = True
            break
        staging = run / (f".{number:04d}.pending")
        try:
            item = convert_tables(load_tables(path, options, cancel), options, cancel)
            write_package(item, staging, options, path, cancel)
            destination = run / f"{number:04d}_{path.stem[:60]}"
            staging.rename(destination)
            items.append(item)
            record = {"source": str(path), "status": "成功", "output": str(destination), "rows": len(item.samples), "features": len(item.features)}
            summary["success"] += 1
        except Cancelled:
            summary["cancelled"] = True
            break
        except Skipped as exc:
            record = {"source": str(path), "status": "跳过", "reason": str(exc)}
            summary["skipped"] += 1
        except Exception as exc:
            record = {"source": str(path), "status": "失败", "reason": str(exc)}
            summary["failed"] += 1
        finally:
            if staging.exists():
                shutil.rmtree(staging)
        records.append(record)
        if progress:
            progress(number, len(files), record)
    if options.merge and items and not summary["cancelled"] and not options.metadata_only:
        staging = run / ".merged.pending"
        try:
            merged = merge_converted(items)
            write_package(merged, staging, options, None, cancel)
            staging.rename(run / "merged")
            summary["merge"] = {"status": "成功", "output": str(run / "merged"), "rows": len(merged.samples), "features": len(merged.features), "partial": summary["failed"] > 0}
        except Cancelled:
            summary["cancelled"] = True
        except Exception as exc:
            summary["merge"] = {"status": "失败", "reason": str(exc)}
        finally:
            if staging.exists():
                shutil.rmtree(staging)
    if cancel is not None and cancel.is_set():
        summary["cancelled"] = True
    summary["not_processed"] = len(files) - len(records)
    (run / "batch_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary
